#!/usr/bin/env python3
"""Parsea las hojas ITEMIZADO del Excel de sobrecostos y extrae, por obra y
centro de costo, el COSTO (NETO) y la última proyección registrada completa.

Uso:
    python3 itemizado_neto_proy.py <Sobrecostos.xlsx> [salida.sql]

Genera un INSERT para `SobrecostosDboard_neto_proy` (ver sql/05_neto_proy.sql) y
valida que la suma de costo neto por obra reconcilie con el total de control de
cada hoja. La taxonomía de columnas varía por hoja, así que se localizan por
texto de encabezado, no por posición:

  * Fila de encabezado = la que contiene "CENTROS DE COSTOS".
  * Columna NETO = la columna "COSTOS" inmediatamente a la izquierda de "IVA".
  * Filas de centro de costo = outline_level 0 dentro del bloque del itemizado
    (corta en la primera fila con nombre vacío, antes del Cuadro de Resumen).
  * Última proyección = la columna PROY más a la derecha COMPLETA (cobertura
    >= 80% de los CC). Excluye "POR GASTAR" y "PROYECCION" sin mes; si la del
    último mes está a medio cargar, cae a "PROY. MES ANTERIOR".
"""
import sys, re, json, datetime
import openpyxl

MES = {'ENE':1,'FEB':2,'MAR':3,'ABR':4,'MAY':5,'JUN':6,'JUL':7,'AGO':8,'SEP':9,
       'OCT':10,'NOV':11,'DIC':12,'ENERO':1,'FEBRERO':2,'MARZO':3,'ABRIL':4,
       'MAYO':5,'JUNIO':6,'JULIO':7,'AGOSTO':8,'SEPTIEMBRE':9,'SEPT':9,
       'OCTUBRE':10,'NOVIEMBRE':11,'DICIEMBRE':12}
ABR_ES = {1:'ENE',2:'FEB',3:'MAR',4:'ABR',5:'MAY',6:'JUN',7:'JUL',8:'AGO',
          9:'SEP',10:'OCT',11:'NOV',12:'DIC'}


def parse_mes(h):
    u = re.sub(r'\s+', ' ', str(h)).strip().upper()
    u = u.replace('PROY', '').replace('PROYECCIÓN', '').replace('PROYECCION', '')
    m = re.search(r'([A-ZÁÉÍÓÚ]{3,10})[\.,\- ]*(\d{4}|\d{2})\b', u)
    if not m or m.group(1) not in MES:
        return None
    y = int(m.group(2)); y = y + 2000 if y < 100 else y
    return datetime.date(y, MES[m.group(1)], 1).isoformat()


def kind(h):
    """('mes', iso) | ('anterior', None) | None para columnas PROY elegibles."""
    u = re.sub(r'\s+', ' ', str(h)).strip().upper()
    if not u.startswith('PROY') or 'GASTAR' in u:
        return None
    p = parse_mes(h)
    if p:
        return ('mes', p)
    if 'MES ANT' in u or 'ANTERIOR' in u:
        return ('anterior', None)
    return None


def obra_of(name):
    m = re.search(r'([A-Z]{2,4})[ _]+(\d{3})', name.upper())
    return f"{m.group(1)}_{m.group(2)}"


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []
    for name in wb.sheetnames:
        if not name.upper().startswith('ITEMIZADO'):
            continue
        ws = wb[name]; obra = obra_of(name)
        hr = name_col = None
        for r in range(1, 21):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v and 'CENTROS DE CO' in str(v).upper():
                    hr, name_col = r, c; break
            if hr:
                break
        iva_col = next(c for c in range(1, ws.max_column + 1)
                       if str(ws.cell(hr, c).value).strip().upper() == 'IVA')
        neto_col = iva_col - 1
        ccs = []
        for r in range(hr + 1, ws.max_row + 1):
            nm = ws.cell(r, name_col).value
            if nm is None or str(nm).strip() == '':
                break
            if ws.row_dimensions[r].outline_level != 0:
                continue
            s = str(nm).strip()
            if re.match(r'^\d{3}\b', s):
                ccs.append((r, s))
        elig = []
        for c in range(1, ws.max_column + 1):
            k = kind(ws.cell(hr, c).value)
            if not k:
                continue
            cov = sum(1 for r, _ in ccs if num(ws.cell(r, c).value) != 0) / len(ccs)
            elig.append((c, k[1], cov))
        full = [e for e in elig if e[2] >= 0.8]
        pc, piso, _ = max(full, key=lambda e: e[0]) if full \
            else max(elig, key=lambda e: (e[2], e[0]))
        label = (f"{ABR_ES[datetime.date.fromisoformat(piso).month]}-"
                 f"{datetime.date.fromisoformat(piso).year}") if piso else 'Mes anterior'
        suma = 0
        for r, s in ccs:
            cc = re.match(r'^(\d{3})', s).group(1)
            neto = num(ws.cell(r, neto_col).value)
            suma += neto
            rows.append(dict(obra=obra, cc_codigo=cc,
                             cc_nombre=re.sub(r'^\d{3}[\s\-]*', '', s)[:80],
                             costo_neto=round(neto, 2),
                             proy_ultima=round(num(ws.cell(r, pc).value), 2),
                             proy_periodo=piso, proy_label=label))
        ftot = num(ws.cell(hr - 1, neto_col).value)
        flag = 'OK' if abs(suma - ftot) < 1 else 'XX'
        print(f"{flag} {obra:8} CCs={len(ccs):2} neto={suma:>16,.0f} (tot {ftot:>16,.0f}) proy={label}")
    return rows


def to_sql(rows):
    def s(x): return "'" + str(x).replace("'", "''") + "'"
    vals = [f"({s(r['obra'])},{s(r['cc_codigo'])},{s(r['cc_nombre'])},"
            f"{r['costo_neto']},{r['proy_ultima']},"
            f"{'null' if r['proy_periodo'] is None else s(r['proy_periodo'])},"
            f"{s(r['proy_label'])})" for r in rows]
    cols = "obra,cc_codigo,cc_nombre,costo_neto,proy_ultima,proy_periodo,proy_label"
    return (f'insert into "SobrecostosDboard_neto_proy"({cols}) values\n'
            + ",\n".join(vals) + ";")


if __name__ == '__main__':
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else 'insert_neto_proy.sql'
    data = parse(src)
    open(out, 'w').write(to_sql(data))
    json.dump(data, open(out.replace('.sql', '.json'), 'w'),
              ensure_ascii=False, indent=0)
    print(f"\nfilas: {len(data)}  ->  {out}")
