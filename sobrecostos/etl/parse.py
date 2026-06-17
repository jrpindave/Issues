import openpyxl, re, json, datetime
F='/root/.claude/uploads/71302bad-fd03-5633-84ee-7ab529a192e2/a4246e1f-Sobrecostos.xlsx'
wb=openpyxl.load_workbook(F, read_only=True, data_only=True)
def norm(s): return re.sub(r'\s+',' ',str(s)).strip() if s is not None else ''
def up(s): return norm(s).upper()
def obra_of(n):
    m=re.search(r'([A-Z]{2,4})[ _]+?(\d{3})', n.upper()); return f"{m.group(1)}_{m.group(2)}" if m else n.upper()
MES={'ENE':1,'ENERO':1,'FEB':2,'FEBRERO':2,'MAR':3,'MARZO':3,'ABR':4,'ABRIL':4,'MAY':5,'MAYO':5,
 'JUN':6,'JUNIO':6,'JUL':7,'JULIO':7,'AGO':8,'AGOSTO':8,'SEP':9,'SEPT':9,'SEPTIEMBRE':9,
 'OCT':10,'OCTUBRE':10,'NOV':11,'NOVIEMBRE':11,'DIC':12,'DICIEMBRE':12}
def parse_periodo(h):
    u=up(h)
    if 'ANTERIOR' in u or 'GASTAR' in u or 'VARIAC' in u: return None
    m=re.search(r'([A-ZÁÉÍÓÚ]{3,10})[\.\- ]*?(\d{4}|\d{2})\b', u.replace('PROY',' ').replace('PROYECCIÓN',' '))
    if m and m.group(1) in MES:
        y=int(m.group(2)); y=y+2000 if y<100 else y
        return datetime.date(y, MES[m.group(1)],1).isoformat()
    return None
def num(v):
    if v is None or v=='': return None
    if isinstance(v,bool): return None
    if isinstance(v,(int,float)): return float(v)
    try: return float(re.sub(r'[^0-9eE+\-.]','',str(v)))
    except: return None

prpc=[]; item=[]; proy=[]; _seen=set()
for name in wb.sheetnames:
    ws=wb[name]; rows=list(ws.iter_rows(values_only=True)); obra=obra_of(name)
    is_prp=up(name).startswith('PRP')
    hidx=None
    for i,r in enumerate(rows[:14]):
        j=' | '.join(up(c) for c in r)
        if is_prp and 'PRESUPUESTO' in j and 'AVANCE' in j: hidx=i; break
        if not is_prp and 'CENTROS DE CO' in j and 'COMPRADO' in j: hidx=i; break
    if hidx is None: print('!! no header',name); continue
    hdr=[up(c) for c in rows[hidx]]
    def find(*keys, start=0, end=None):
        end=len(hdr) if end is None else end
        for idx in range(start,end):
            if all(k in hdr[idx] for k in keys): return idx
        return None
    if is_prp:
        c_pres=find('PRESUPUESTO'); c_comb=c_pres-2; c_ctac=c_pres-1
        c_prog=find('PROGRAMAC',start=c_pres); c_av=find('AVANCE',start=c_pres)
        c_rc=find('REAL CONTABLE',start=c_pres) or find('REAL CONTA',start=c_pres)
        c_ro=find('REAL OBRA',start=c_pres); c_tras=find('TRASPASO',start=c_pres)
        cur=None
        for r in rows[hidx+1:]:
            comb=norm(r[c_comb]) if c_comb<len(r) else ''
            ctac=norm(r[c_ctac]) if c_ctac<len(r) else ''
            if not comb: continue
            es=(ctac=='')
            if es: cur=comb
            g=lambda i:(num(r[i]) if i is not None and i<len(r) else None)
            prpc.append(dict(obra=obra,centro_costo=cur,etiqueta=comb,cuenta_codigo=ctac,es_centro_costo=es,
                presupuesto=g(c_pres),programacion=g(c_prog),avance=g(c_av),
                real_contable=g(c_rc),real_obra=g(c_ro),traspaso=g(c_tras)))
    else:
        ci=find('CENTROS DE CO')
        c_clase=find('CLASE'); c_tipo=find('TIPO'); c_cc=(0 if hdr and hdr[0]=='CC' else find('CC',end=ci))
        c_rc=find('CONTABL',start=ci+5); c_ro=find('REAL OBRA',start=ci)
        proy_cols=[(i,parse_periodo(rows[hidx][i])) for i,h in enumerate(hdr) if h.startswith('PROY') or 'PROYECC' in h]
        proy_cols=[(i,p) for i,p in proy_cols if p]
        for r in rows[hidx+1:]:
            it=norm(r[ci]) if ci<len(r) else ''
            if not re.match(r'^\d{3}\b', it): continue   # solo lineas de costo (codigo CC)
            tipo=norm(r[c_tipo]) if c_tipo is not None and c_tipo<len(r) else ''
            clase=norm(r[c_clase]) if c_clase is not None and c_clase<len(r) else ''
            cc_code=re.match(r'^(\d{3})',it).group(1)
            key=(obra,cc_code)
            nivel='CC' if key not in _seen else 'CTA'
            _seen.add(key)
            g=lambda i:(num(r[i]) if i is not None and i<len(r) else None)
            item.append(dict(obra=obra,clase=clase,tipo=tipo,nivel=nivel,
                cc_codigo=cc_code,
                cc=norm(r[c_cc]) if c_cc is not None and c_cc<len(r) else '',
                item=it,pct=g(ci+1),costo_neto=g(ci+2),iva=g(ci+3),costo_con_iva=g(ci+4),
                comprado=g(ci+5),recepcionado=g(ci+6),real_contable=g(c_rc),real_obra=g(c_ro)))
            for pc,per in proy_cols:
                v=g(pc)
                if v is None: continue
                proy.append(dict(obra=obra,item=it,cc_codigo=cc_code,col_idx=pc,periodo=per,monto=v))

json.dump(prpc,open('sobrecostos_build/prpc.json','w'))
json.dump(item,open('sobrecostos_build/item.json','w'))
json.dump(proy,open('sobrecostos_build/proy.json','w'))
from collections import defaultdict
print('TOTALES prpc:%d item:%d proy:%d'%(len(prpc),len(item),len(proy)))
ag=defaultdict(lambda:[0,0.0,0.0])
for r in item: ag[r['obra']][0]+=1; ag[r['obra']][1]+=(r['costo_neto'] or 0); ag[r['obra']][2]+=(r['costo_con_iva'] or 0)
print('\nITEMIZADO obra | filas | sum costo_neto | sum costo_con_iva')
for o in sorted(ag): print(f'  {o:8} {ag[o][0]:4}  neto={ag[o][1]:>18,.0f}  c_iva={ag[o][2]:>18,.0f}')
# dup check SP vs MU numeric
sp=[r for r in item if r['obra']=='SP_296']; mu=[r for r in item if r['obra']=='MU_293']
same = len(sp)==len(mu) and all(abs((a['costo_con_iva'] or 0)-(b['costo_con_iva'] or 0))<1 for a,b in zip(sp,mu))
print('\nITEMIZADO SP_296 == MU_293 (numericamente identicas)?:', same)
